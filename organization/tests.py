import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string as CI

from config.testutils import BaseTest
from enrollment.models import EnrollmentEntry, EnrollmentRecord, RecordSource, RecordStatus
from enrollment.schema import AGGREGATES, LEVELS
from enrollment.services import get_or_create_record, save_values
from .importer import apply_import, parse_workbook, reconcile
from .models import District, School


def make_workbook(rows, banner=True, corrupt_header=False):
    """rows: list of dicts(school_id, name, district, values={level:(m,f)}, totals=None)"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Public Schools"
    for key, label, _s, col in LEVELS:
        ws.cell(7, CI(col)).value = "WRONG" if (corrupt_header and key == "g3") else label
    r = 9
    for n, row in enumerate(rows, 1):
        ws.cell(r, 1).value = n
        ws.cell(r, 2).value = row["school_id"]
        ws.cell(r, 3).value = row["name"]
        ws.cell(r, 4).value = row.get("district", "Atok")
        ws.cell(r, 5).value = "Public"
        for key, _l, _s, col in LEVELS:
            m, f = row.get("values", {}).get(key, (None, None))
            ws.cell(r, CI(col)).value = m
            ws.cell(r, CI(col) + 1).value = f
        for key, (_l, _p, col) in AGGREGATES.items():          # stand-in for Excel's cached formula results
            if row.get("totals") and key in row["totals"]:
                for i, v in enumerate(row["totals"][key]):
                    ws.cell(r, CI(col) + i).value = v
        r += 1
    if banner:
        ws.cell(r, 3).value = "ATOK DISTRICT"
        ws.cell(r + 1, 3).value = "TOTAL ENROLLMENT-SDO-BENGUET"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


ROWS = [
    {"school_id": 111111, "name": "Alpha ES", "district": "Atok", "values": {"kinder": (6, 4), "g1": (10, 7)}},
    {"school_id": 222222, "name": "Beta NHS", "district": "Bakun", "values": {"g7": (1, 2)}},
]


class ParseTests(TestCase):
    def test_parses_rows_ignores_banners(self):
        p = parse_workbook(make_workbook(ROWS))
        self.assertEqual(p.errors, [])
        self.assertEqual([r.school_id for r in p.rows], ["111111", "222222"])
        self.assertEqual(p.rows[0].values["kinder"], (6, 4))
        self.assertIn("ATOK DISTRICT", p.banners)

    def test_header_drift_rejected(self):
        p = parse_workbook(make_workbook(ROWS, corrupt_header=True))
        self.assertTrue(any("Header mismatch" in e for e in p.errors))

    def test_row_errors_reported_not_silently_dropped(self):
        bad = ROWS + [
            {"school_id": 111111, "name": "Dup ID", "district": "Atok"},
            {"school_id": 333333, "name": "No District", "district": ""},
            {"school_id": 444444, "name": "Bad Value", "district": "Atok", "values": {"g1": ("abc", 3)}},
            {"school_id": 555555, "name": "Negative", "district": "Atok", "values": {"g1": (-2, 3)}},
            {"school_id": "12AB", "name": "Bad ID", "district": "Atok"},
        ]
        p = parse_workbook(make_workbook(bad))
        text = " | ".join(p.errors)
        for needle in ("duplicate school ID 111111", "no district", "expected a whole number", "invalid school ID"):
            self.assertIn(needle, text)

    def test_whitespace_cell_is_warning_not_error(self):
        rows = [{"school_id": 1, "name": "Spacey", "district": "Atok", "values": {"g11_sports": ("  ", 2)}}]
        p = parse_workbook(make_workbook(rows))
        self.assertEqual(p.errors, [])
        self.assertTrue(p.warnings)
        self.assertEqual(p.rows[0].values["g11_sports"], (None, 2))

    def test_duplicate_name_in_same_district_is_warning(self):
        rows = [{"school_id": 1, "name": "Same", "district": "Atok"}, {"school_id": 2, "name": "same", "district": "Atok"},
                {"school_id": 3, "name": "Same", "district": "Bakun"}]
        p = parse_workbook(make_workbook(rows))
        self.assertEqual(p.errors, [])
        self.assertEqual(len([w for w in p.warnings if "appears twice" in w]), 1)


class ApplyTests(TestCase):
    def test_dry_run_writes_nothing(self):
        rep = apply_import(parse_workbook(make_workbook(ROWS)), with_enrollment=True, dry_run=True)
        self.assertEqual((rep["schools_created"], rep["committed"]), (2, False))
        self.assertEqual((School.objects.count(), District.objects.count(), EnrollmentRecord.objects.count()), (0, 0, 0))

    def test_commit_creates_and_preserves_ids_and_districts(self):
        rep = apply_import(parse_workbook(make_workbook(ROWS)), with_enrollment=True)
        self.assertTrue(rep["committed"])
        self.assertEqual(School.objects.get(school_id="222222").district.name, "Bakun")
        self.assertEqual(School.objects.get(school_id="111111").name, "Alpha ES")
        self.assertEqual(District.objects.count(), 2)
        rec = EnrollmentRecord.objects.get(school__school_id="111111")
        self.assertEqual((rec.status, rec.source), (RecordStatus.DRAFT, RecordSource.EXCEL_IMPORT))
        self.assertEqual(EnrollmentEntry.objects.get(record=rec, level="g1").male, 10)

    def test_reimport_is_idempotent(self):
        for _ in range(3):
            apply_import(parse_workbook(make_workbook(ROWS)), with_enrollment=True)
        self.assertEqual((School.objects.count(), District.objects.count(), EnrollmentRecord.objects.count()), (2, 2, 2))
        self.assertEqual(EnrollmentEntry.objects.count(), 2 * len(LEVELS))

    def test_errors_cause_full_rollback_nothing_written(self):
        bad = ROWS + [{"school_id": 111111, "name": "Dup", "district": "Atok"}]
        rep = apply_import(parse_workbook(make_workbook(bad)), with_enrollment=True)
        self.assertTrue(rep["errors"])
        self.assertFalse(rep["committed"])
        self.assertEqual(School.objects.count(), 0)

    def test_failure_midway_rolls_back_transaction(self):
        from unittest import mock
        with mock.patch("organization.importer._baseline", side_effect=RuntimeError("boom")):
            with self.assertRaises(RuntimeError):
                apply_import(parse_workbook(make_workbook(ROWS)), with_enrollment=True)
        self.assertEqual((School.objects.count(), District.objects.count()), (0, 0))

    def test_portal_entered_and_submitted_data_never_overwritten(self):
        apply_import(parse_workbook(make_workbook(ROWS)), with_enrollment=True)
        school = School.objects.get(school_id="111111")
        rec = get_or_create_record(school)
        save_values(rec, {"g1": (99, 98)}, None)             # user edit -> source becomes manual
        changed = [dict(ROWS[0], values={"g1": (1, 1)}), ROWS[1]]
        rep = apply_import(parse_workbook(make_workbook(changed)), with_enrollment=True)
        self.assertEqual(rep["records_preserved"], 1)
        self.assertEqual(EnrollmentEntry.objects.get(record=rec, level="g1").male, 99)

    def test_school_moved_between_districts_is_reported(self):
        apply_import(parse_workbook(make_workbook(ROWS)))
        moved = [dict(ROWS[0], district="Bakun"), ROWS[1]]
        rep = apply_import(parse_workbook(make_workbook(moved)))
        self.assertEqual(len(rep["schools_moved"]), 1)
        self.assertEqual(School.objects.get(school_id="111111").district.name, "Bakun")

    def test_reconcile_detects_typed_total_that_disagrees(self):
        rows = [dict(ROWS[0], totals={"es": (6 + 10, 4 + 7, 28)}),          # T typed as 28; M+F is 27
                dict(ROWS[1], totals={"jhs": (1, 2, 3)})]
        parsed = parse_workbook(make_workbook(rows))
        apply_import(parsed, with_enrollment=True)
        rec = reconcile(parsed)
        self.assertEqual(rec["rows_checked"], 2)
        self.assertEqual(len(rec["mismatches"]), 1)
        self.assertIn("ELEMENTARY TOTAL", rec["mismatches"][0])


class CommandAndViewTests(BaseTest):
    def test_command_rejects_bad_file(self):
        with self.assertRaises(CommandError):
            call_command("import_master", "/nonexistent.xlsx")

    def test_import_view_is_system_admin_only(self):
        for u in (self.u1, self.admin1):
            self.login(u)
            self.assertEqual(self.client.get(reverse("manage:import")).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(reverse("manage:import")).status_code, 302)

    def test_import_view_dry_run_then_commit(self):
        self.login(self.sysadmin)
        data = make_workbook(ROWS).getvalue()
        before = School.objects.count()
        r = self.client.post(reverse("manage:import"), {"workbook": SimpleUploadedFile("w.xlsx", data), "dry_run": "on"})
        self.assertContains(r, "Dry run")
        self.assertEqual(School.objects.count(), before)
        r = self.client.post(reverse("manage:import"), {"workbook": SimpleUploadedFile("w.xlsx", data)})
        self.assertContains(r, "Import committed")
        self.assertEqual(School.objects.count(), before + 2)

    def test_import_view_rejects_non_xlsx_and_garbage(self):
        self.login(self.sysadmin)
        r = self.client.post(reverse("manage:import"), {"workbook": SimpleUploadedFile("evil.exe", b"MZ")})
        self.assertContains(r, "Upload an .xlsx")
        r = self.client.post(reverse("manage:import"), {"workbook": SimpleUploadedFile("x.xlsx", b"not a zip")})
        self.assertContains(r, "could not be read")

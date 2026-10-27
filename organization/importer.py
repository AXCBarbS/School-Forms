"""Import the district/school master list (and optionally enrollment baselines) from the
consolidated SF4 workbook. The source file is opened read-only and never modified."""
from dataclasses import dataclass, field

import openpyxl
from django.conf import settings
from django.db import transaction
from openpyxl.utils import column_index_from_string as CI

from enrollment.models import EnrollmentEntry, EnrollmentRecord, RecordSource, RecordStatus
from enrollment.schema import AGGREGATES, LEVELS, LEVEL_KEYS
from enrollment.services import compute_totals
from .models import District, School

HEADER_ROW, FIRST_DATA_ROW = 7, 9
COL = {"no": 1, "school_id": 2, "name": 3, "district": 4, "sector": 5}
BANNER_MARKERS = ("DISTRICT", "TOTAL ENROLLMENT")


class DryRun(Exception):
    pass


@dataclass
class ParsedRow:
    row: int
    order: int
    school_id: str
    name: str
    district: str
    sector: str
    values: dict
    cached: dict          # {aggregate_key: (M, F, T)} from the workbook's saved formula results


@dataclass
class Parsed:
    rows: list = field(default_factory=list)
    banners: dict = field(default_factory=dict)       # {banner text: cached K-12 total}
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)


def _int_or_none(v, where, parsed, input_cell=True):
    if v is None or (isinstance(v, str) and not v.strip()):
        if isinstance(v, str):
            parsed.warnings.append(f"{where}: whitespace-only text treated as blank")
        return None
    if isinstance(v, bool):
        parsed.errors.append(f"{where}: invalid value {v!r}")
        return None
    if isinstance(v, (int, float)) and float(v) == int(v) and v >= 0:
        return int(v)
    parsed.errors.append(f"{where}: expected a whole number >= 0, found {v!r}")
    return None


def parse_workbook(source):
    parsed = Parsed()
    wb_v = openpyxl.load_workbook(source, data_only=True)
    if hasattr(source, "seek"):
        source.seek(0)
    wb_f = openpyxl.load_workbook(source)
    ws_v, ws_f = wb_v.worksheets[0], wb_f.worksheets[0]
    # structure check: level labels must sit where the schema expects them
    for key, label, _sec, col in LEVELS:
        found = str(ws_v.cell(HEADER_ROW, CI(col)).value or "").strip()
        if found.lower() != label.lower():
            parsed.errors.append(f"Header mismatch at {col}{HEADER_ROW}: expected '{label}', found '{found}'")
    if parsed.errors:
        return parsed
    seen_ids = {}
    seen_names = {}
    for r in range(FIRST_DATA_ROW, ws_v.max_row + 1):
        raw_id, name = ws_v.cell(r, COL["school_id"]).value, ws_v.cell(r, COL["name"]).value
        label = str(name or "").strip()
        if raw_id in (None, ""):
            if any(m in label.upper() for m in BANNER_MARKERS):
                parsed.banners[label] = ws_v.cell(r, CI("CB")).value
            elif label:
                parsed.errors.append(f"Row {r}: '{label}' has no school ID")
            continue
        sid = str(raw_id).strip()
        if sid.endswith(".0"):
            sid = sid[:-2]
        if not sid.isdigit():
            parsed.errors.append(f"Row {r}: invalid school ID {raw_id!r}")
            continue
        district = str(ws_v.cell(r, COL["district"]).value or "").strip()
        if not label:
            parsed.errors.append(f"Row {r}: school {sid} has no name")
        if not district:
            parsed.errors.append(f"Row {r}: school {sid} has no district")
        if sid in seen_ids:
            parsed.errors.append(f"Row {r}: duplicate school ID {sid} (first at row {seen_ids[sid]})")
            continue
        seen_ids[sid] = r
        key = (district.lower(), label.lower())
        if key in seen_names:
            parsed.warnings.append(f"Row {r}: '{label}' appears twice in district {district} (rows {seen_names[key]}, {r})")
        seen_names[key] = r
        values = {}
        for k, lab, _s, col in LEVELS:
            c = CI(col)
            for off, who in ((0, "Male"), (1, "Female")):
                if isinstance(ws_f.cell(r, c + off).value, str) and str(ws_f.cell(r, c + off).value).startswith("="):
                    parsed.warnings.append(f"Row {r} {lab} {who}: formula in an input cell; cached value used")
            m = _int_or_none(ws_v.cell(r, c).value, f"Row {r} {lab} Male", parsed)
            f = _int_or_none(ws_v.cell(r, c + 1).value, f"Row {r} {lab} Female", parsed)
            values[k] = (m, f)
        cached = {}
        for key2, (_l, _p, col) in AGGREGATES.items():
            c = CI(col)
            cached[key2] = tuple(ws_v.cell(r, c + i).value for i in range(3))
        parsed.rows.append(ParsedRow(r, len(parsed.rows) + 1, sid, label, district,
                                     str(ws_v.cell(r, COL["sector"]).value or "Public").strip(), values, cached))
    if not parsed.rows and not parsed.errors:
        parsed.errors.append("No school rows were found.")
    return parsed


def apply_import(parsed, with_enrollment=False, dry_run=False, user=None):
    """All-or-nothing. Returns a report dict. With errors in `parsed` nothing is written."""
    rep = {"districts_created": 0, "schools_created": 0, "schools_updated": 0, "schools_unchanged": 0,
           "schools_moved": [], "records_created": 0, "records_updated": 0, "records_preserved": 0,
           "errors": list(parsed.errors), "warnings": list(parsed.warnings), "committed": False,
           "rows_read": len(parsed.rows)}
    if parsed.errors:
        return rep
    try:
        with transaction.atomic():
            districts = {}
            for pr in parsed.rows:
                if pr.district not in districts:
                    d, created = District.objects.get_or_create(name=pr.district)
                    districts[pr.district] = d
                    rep["districts_created"] += created
            for pr in parsed.rows:
                school = School.objects.filter(school_id=pr.school_id).first()
                d = districts[pr.district]
                if school is None:
                    school = School.objects.create(school_id=pr.school_id, name=pr.name, district=d,
                                                   sector=pr.sector, sort_order=pr.order)
                    rep["schools_created"] += 1
                else:
                    changed = (school.name, school.district_id, school.sector, school.sort_order) != \
                              (pr.name, d.id, pr.sector, pr.order)
                    if school.district_id != d.id:
                        rep["schools_moved"].append(f"{pr.school_id} {pr.name}: {school.district.name} -> {d.name}")
                    if changed:
                        school.name, school.district, school.sector, school.sort_order = pr.name, d, pr.sector, pr.order
                        school.save()
                        rep["schools_updated"] += 1
                    else:
                        rep["schools_unchanged"] += 1
                if with_enrollment:
                    _baseline(school, pr, rep, user)
            if dry_run:
                raise DryRun()
            rep["committed"] = True
    except DryRun:
        pass
    return rep


def _baseline(school, pr, rep, user):
    year = settings.CURRENT_SCHOOL_YEAR
    record = EnrollmentRecord.objects.filter(school=school, school_year=year).first()
    if record is None:
        record = EnrollmentRecord.objects.create(school=school, school_year=year, source=RecordSource.EXCEL_IMPORT,
                                                 status=RecordStatus.DRAFT, created_by=user, updated_by=user)
        rep["records_created"] += 1
    elif record.source != RecordSource.EXCEL_IMPORT or record.status == RecordStatus.SUBMITTED:
        rep["records_preserved"] += 1       # never overwrite data entered or submitted in the portal
        return
    else:
        rep["records_updated"] += 1
    existing = {e.level: e for e in record.entries.all()}
    for k in LEVEL_KEYS:
        e = existing.get(k) or EnrollmentEntry(record=record, level=k)
        e.male, e.female = pr.values[k]
        e.save()


def reconcile(parsed):
    """Compare what the database now holds with the workbook's saved results.
    Returns {'rows_checked', 'mismatches': [...], 'district_totals': [...]}"""
    year = settings.CURRENT_SCHOOL_YEAR
    mismatches, per_district = [], {}
    by_id = {r.school_id: r for r in parsed.rows}
    records = EnrollmentRecord.objects.filter(school__school_id__in=by_id, school_year=year) \
        .select_related("school__district").prefetch_related("entries")
    seen = set()
    for rec in records:
        pr = by_id[rec.school.school_id]
        seen.add(pr.school_id)
        vals = {e.level: (e.male, e.female) for e in rec.entries.all()}
        tot = compute_totals(vals)
        d = per_district.setdefault(rec.school.district.name, {"db": 0, "wb": 0})
        d["db"] += tot["k12"][2]
        for key in AGGREGATES:
            cached = pr.cached[key]
            if None in cached:
                continue
            if tuple(tot[key]) != tuple(cached):
                mismatches.append(f"Row {pr.row} {pr.name} [{pr.district}] {AGGREGATES[key][0]}: "
                                  f"workbook M/F/T={tuple(cached)} but computed={tuple(tot[key])}")
        d["wb"] += pr.cached["k12"][2] or 0
    for pr in parsed.rows:
        if pr.school_id not in seen:
            mismatches.append(f"Row {pr.row} {pr.name}: no enrollment record in database")
    return {"rows_checked": len(seen), "mismatches": mismatches,
            "district_totals": sorted(per_district.items())}

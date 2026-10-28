
from datetime import datetime

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import EnrollmentRecord
from .schema import (
    AGGREGATE_LABEL,
    EXPORT_ORDER,
    LEVEL_LABEL,
)
from .services import compute_totals, record_state, STATE_LABELS


MOVEMENT_CATEGORIES = [
    ("dropout", "NLPA / Dropout"),
    ("transferred_out", "Transferred Out"),
    ("transferred_in", "Transferred In"),
]

MOVEMENT_PERIODS = [
    ("previous", "A - Cumulative as of Previous Month"),
    ("monthly", "B - For the Month"),
    ("cumulative", "A+B - Cumulative as of End of the Month"),
]


def movement_totals(record):
    """Return M, F, T values for all SF4 movement reporting periods."""
    results = {}

    saved_entries = (
        {
            entry.category: entry
            for entry in record.movement_entries.all()
        }
        if record
        else {}
    )

    for category, _label in MOVEMENT_CATEGORIES:
        entry = saved_entries.get(category)

        previous_male = (
            entry.previous_male or 0 if entry else 0
        )
        previous_female = (
            entry.previous_female or 0 if entry else 0
        )
        monthly_male = (
            entry.monthly_male or 0 if entry else 0
        )
        monthly_female = (
            entry.monthly_female or 0 if entry else 0
        )

        cumulative_male = previous_male + monthly_male
        cumulative_female = previous_female + monthly_female

        results[(category, "previous")] = (
            previous_male,
            previous_female,
            previous_male + previous_female,
        )

        results[(category, "monthly")] = (
            monthly_male,
            monthly_female,
            monthly_male + monthly_female,
        )

        results[(category, "cumulative")] = (
            cumulative_male,
            cumulative_female,
            cumulative_male + cumulative_female,
        )

    return results


def build_export(schools):
    """
    Build an Excel workbook with one row per school.

    Existing enrollment columns are preserved. SF4 movement columns
    are appended after the existing enrollment columns.
    """
    year = settings.CURRENT_SCHOOL_YEAR

    wb = Workbook()
    ws = wb.active
    ws.title = "Public Schools"

    # Existing enrollment headings.
    fixed_headers = [
        "No.",
        "School ID",
        "School Name",
        "District",
        "Sector",
        "Status",
    ]

    head = list(fixed_headers)
    sub = [""] * len(fixed_headers)

    for kind, key in EXPORT_ORDER:
        label = (
            LEVEL_LABEL[key]
            if kind == "level"
            else AGGREGATE_LABEL[key]
        )

        head.extend([label] * 3)
        sub.extend(["M", "F", "T"])

    # New SF4 movement columns: 3 categories x 3 periods x 3 values.
    movement_start_col = len(head) + 1

    for _category, category_label in MOVEMENT_CATEGORIES:
        for _period, period_label in MOVEMENT_PERIODS:
            head.extend([category_label] * 3)
            sub.extend(["M", "F", "T"])

    ws.append([f"SF4 Enrollment {year}"])
    ws.append(head)
    ws.append(sub)

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDE7F3")
    movement_fill = PatternFill("solid", fgColor="E2F0D9")

    # Style the title and two header rows.
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(head),
    )

    for row in ws.iter_rows(min_row=1, max_row=3):
        for cell in row:
            cell.font = bold
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    # Merge existing enrollment group headings.
    column = len(fixed_headers) + 1

    for _kind, _key in EXPORT_ORDER:
        ws.merge_cells(
            start_row=2,
            start_column=column,
            end_row=2,
            end_column=column + 2,
        )
        column += 3

    # Merge each SF4 category and reporting-period heading.
    column = movement_start_col

    for _category, category_label in MOVEMENT_CATEGORIES:
        category_start = column

        for _period, period_label in MOVEMENT_PERIODS:
            ws.merge_cells(
                start_row=2,
                start_column=column,
                end_row=2,
                end_column=column + 2,
            )

            ws.cell(row=2, column=column).value = (
                f"{category_label}\n{period_label}"
            )

            for offset in range(3):
                ws.cell(
                    row=2,
                    column=column + offset,
                ).fill = movement_fill

            column += 3

        # The category label is included in each merged period heading.
        # This keeps every column group identifiable in Excel.

    records = {
        record.school_id: record
        for record in EnrollmentRecord.objects.filter(
            school__in=schools,
            school_year=year,
        ).prefetch_related("entries", "movement_entries")
    }

    grand_enrollment = {}
    grand_movement = {}

    ordered_schools = schools.order_by(
        "district__name",
        "sort_order",
        "name",
    )

    for number, school in enumerate(ordered_schools, 1):
        record = records.get(school.id)

        enrollment_values = (
            {
                entry.level: (entry.male, entry.female)
                for entry in record.entries.all()
            }
            if record
            else {}
        )

        totals = compute_totals(enrollment_values)

        row = [
            number,
            school.school_id,
            school.name,
            school.district.name,
            school.sector,
            STATE_LABELS[record_state(record)][0],
        ]

        # Existing enrollment data.
        for kind, key in EXPORT_ORDER:
            male, female, total = totals[key]
            row.extend([male, female, total])

            grand = grand_enrollment.setdefault(key, [0, 0, 0])
            grand[0] += male
            grand[1] += female
            grand[2] += total

        # New movement data.
        movements = movement_totals(record)

        for category, _category_label in MOVEMENT_CATEGORIES:
            for period, _period_label in MOVEMENT_PERIODS:
                values = movements[(category, period)]
                row.extend(values)

                grand_key = (category, period)
                grand = grand_movement.setdefault(
                    grand_key,
                    [0, 0, 0],
                )
                grand[0] += values[0]
                grand[1] += values[1]
                grand[2] += values[2]

        ws.append(row)

    # Grand-total row.
    total_row = ["", "", "TOTAL", "", "", ""]

    for _kind, key in EXPORT_ORDER:
        total_row.extend(grand_enrollment.get(key, [0, 0, 0]))

    for category, _category_label in MOVEMENT_CATEGORIES:
        for period, _period_label in MOVEMENT_PERIODS:
            total_row.extend(
                grand_movement.get(
                    (category, period),
                    [0, 0, 0],
                )
            )

    ws.append(total_row)

    for cell in ws[ws.max_row]:
        cell.font = bold
        cell.fill = header_fill

    # Worksheet layout.
    ws.freeze_panes = "G4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(head))}{ws.max_row - 1}"

    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 38
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 16

    for index in range(7, len(head) + 1):
        ws.column_dimensions[get_column_letter(index)].width = 9

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 58
    ws.row_dimensions[3].height = 22

    ws.append([])
    ws.append([
        f"Exported {datetime.now():%Y-%m-%d %H:%M}. "
        "Confidential: contains school enrollment data."
    ])

    return wb

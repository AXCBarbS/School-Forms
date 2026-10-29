import io

from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import Approval, User
from audit.models import AuditLog
from config.testutils import PW, BaseTest, make_user
from .models import EnrollmentEntry, EnrollmentRecord, RecordSource, RecordStatus
from .schema import LEVEL_KEYS
from .services import compute_totals, get_or_create_record, save_values, submit_problems

# Adoyunan Elementary School (Atok) row from the workbook: Kinder..Grade 6 M/F
ATOK_ES = {"kinder": (6, 4), "g1": (10, 7), "g2": (5, 6), "g3": (10, 4), "g4": (5, 7), "g5": (4, 10), "g6": (13, 10)}


def form_data(values=None, **extra):
    data = {}
    for k in LEVEL_KEYS:
        m, f = (values or {}).get(k, (None, None))
        data[f"{k}_m"] = "" if m is None else m
        data[f"{k}_f"] = "" if f is None else f
    data.update(extra)
    return data


class CalculationTests(BaseTest):
    def test_matches_workbook_formulas_for_known_school(self):
        t = compute_totals(ATOK_ES)
        self.assertEqual(t["kinder"], (6, 4, 10))                  # H = SUM(F:G)
        self.assertEqual(t["es"], (53, 48, 101))                   # AD/AE/AF
        self.assertEqual(t["k12"], (53, 48, 101))
        self.assertEqual(t["jhs"], (0, 0, 0))

    def test_secondary_aggregates(self):
        vals = {"g7": (10, 12), "sned_jhs": (1, 1), "g11_acad": (22, 27), "g11_tvl": (40, 7), "g12_acad": (9, 31),
                "g12_tvl": (24, 23)}
        t = compute_totals(vals)
        self.assertEqual(t["jhs"], (11, 13, 24))
        self.assertEqual(t["g11"], (62, 34, 96))                   # Twin Peaks NHS Grade 11 (22+40, 27+7)
        self.assertEqual(t["g12"], (33, 54, 87))
        self.assertEqual(t["shs"], (95, 88, 183))
        self.assertEqual(t["k12"], (106, 101, 207))

    def test_blank_counts_as_zero(self):
        self.assertEqual(compute_totals({"g1": (None, 5)})["g1"], (0, 5, 5))

    def test_submit_problems(self):
        empty = {k: (None, None) for k in LEVEL_KEYS}
        self.assertTrue(submit_problems(empty))
        half = dict(empty, g1=(5, None))
        self.assertTrue(any("Grade 1" in p for p in submit_problems(half)))
        ok = dict(empty, g1=(5, 0))
        self.assertEqual(submit_problems(ok), [])


class DashboardAccessTests(BaseTest):
    def test_user_sees_only_own_district_schools(self):
        self.login(self.u1)
        r = self.client.get(reverse("enrollment:dashboard"))
        self.assertContains(r, "Adoyunan ES")
        self.assertContains(r, "Bilis ES")
        self.assertNotContains(r, "Ampusongan ES")
        self.assertContains(r, "Atok District")

    def test_district_query_param_cannot_widen_scope(self):
        self.login(self.u1)
        r = self.client.get(reverse("enrollment:dashboard"), {"district": self.d2.pk})
        self.assertNotContains(r, "Ampusongan ES")

    def test_search_and_status_filter(self):
        self.login(self.u1)
        r = self.client.get(reverse("enrollment:dashboard"), {"q": "adoy"})
        self.assertContains(r, "Adoyunan ES")
        self.assertNotContains(r, "Bilis ES")
        r = self.client.get(reverse("enrollment:dashboard"), {"q": "100002"})
        self.assertContains(r, "Bilis ES")
        r = self.client.get(reverse("enrollment:dashboard"), {"state": "completed"})
        self.assertNotContains(r, "Adoyunan ES")

    def test_status_labels(self):
        rec = get_or_create_record(self.s1a, self.u1)
        rec.source = RecordSource.EXCEL_IMPORT
        rec.save()
        self.login(self.u1)
        self.assertContains(self.client.get(reverse("enrollment:dashboard")), "Pre-filled")
        save_values(rec, ATOK_ES, self.u1)
        self.assertContains(self.client.get(reverse("enrollment:dashboard")), "In Progress")

    def test_system_admin_sees_all_and_can_filter(self):
        self.login(self.sysadmin)
        r = self.client.get(reverse("enrollment:dashboard"))
        self.assertContains(r, "Adoyunan ES")
        self.assertContains(r, "Ampusongan ES")
        r = self.client.get(reverse("enrollment:dashboard"), {"district": self.d2.pk})
        self.assertNotContains(r, "Adoyunan ES")


class CrossDistrictIsolationTests(BaseTest):
    """User of Atok must not read, write, submit, reopen or export Bakun data by any route."""

    def setUp(self):
        super().setUp()
        self.rec2 = get_or_create_record(self.s2a, self.u2)
        save_values(self.rec2, {"g1": (7, 8)}, self.u2)
        self.login(self.u1)

    def test_all_school_routes_404_for_other_district(self):
        for name in ("enrollment:entry", "enrollment:review", "enrollment:submit"):
            r = self.client.get(reverse(name, args=[self.s2a.school_id]))
            self.assertEqual(r.status_code, 404, name)
        for name in ("enrollment:entry", "enrollment:submit", "enrollment:reopen"):
            r = self.client.post(reverse(name, args=[self.s2a.school_id]), form_data({"g1": (99, 99)}))
            self.assertEqual(r.status_code, 404, name)

    def test_other_district_data_unchanged_after_attack(self):
        self.client.post(reverse("enrollment:entry", args=[self.s2a.school_id]), form_data({"g1": (99, 99)}))
        e = EnrollmentEntry.objects.get(record=self.rec2, level="g1")
        self.assertEqual((e.male, e.female), (7, 8))

    def test_nonexistent_and_other_ids_look_identical(self):
        a = self.client.get(reverse("enrollment:entry", args=["999999"]))
        b = self.client.get(reverse("enrollment:entry", args=[self.s2a.school_id]))
        self.assertEqual(a.status_code, b.status_code)

    def test_export_only_contains_own_district(self):
        r = self.client.get(reverse("enrollment:export"), {"district": self.d2.pk})
        ws = load_workbook(io.BytesIO(r.content)).active
        names = [row[2] for row in ws.iter_rows(min_row=4, values_only=True)]
        self.assertIn("Adoyunan ES", names)
        self.assertNotIn("Ampusongan ES", names)

    def test_inactive_school_not_reachable(self):
        self.s1b.is_active = False
        self.s1b.save()
        self.assertEqual(self.client.get(reverse("enrollment:entry", args=[self.s1b.school_id])).status_code, 404)

    def test_user_without_district_sees_nothing(self):
        from django.db import connection
        sysu = self.sysadmin
        self.login(sysu)
        # a system admin always has access; a (hypothetical) user with no district gets none
        from .access import accessible_schools
        ghost = User(email="g@deped.gov.ph", full_name="g", role="user", district=None)
        self.assertEqual(accessible_schools(ghost).count(), 0)


class DataEntryTests(BaseTest):
    def setUp(self):
        super().setUp()
        self.login(self.u1)
        self.url = reverse("enrollment:entry", args=[self.s1a.school_id])

    def test_get_renders_all_levels(self):
        r = self.client.get(self.url)
        self.assertEqual(r.status_code, 200)
        for k in LEVEL_KEYS:
            self.assertContains(r, f'name="{k}_m"')
            self.assertContains(r, f'name="{k}_f"')

    def test_save_draft_persists_and_resumes_after_relogin(self):
        r = self.client.post(self.url, form_data(ATOK_ES, action="save"))
        self.assertRedirects(r, self.url)
        self.client.post(reverse("accounts:logout"))
        self.login(self.u1)
        r = self.client.get(self.url)
        self.assertContains(r, 'value="13"')          # Grade 6 male survives
        rec = EnrollmentRecord.objects.get(school=self.s1a)
        self.assertEqual(rec.status, RecordStatus.DRAFT)
        self.assertEqual(rec.source, RecordSource.MANUAL)
        self.assertEqual(rec.updated_by, self.u1)

    def test_invalid_values_rejected_and_nothing_saved(self):
        for bad in ("-1", "abc", "1.5", "999999"):
            r = self.client.post(self.url, form_data({"g1": (5, 5)}, g2_m=bad))
            self.assertEqual(r.status_code, 200, bad)
        rec = EnrollmentRecord.objects.get(school=self.s1a)
        self.assertEqual(EnrollmentEntry.objects.filter(record=rec, male__isnull=False).count(), 0)

    def test_save_and_review_goes_to_review_with_calculated_totals(self):
        r = self.client.post(self.url, form_data(ATOK_ES, action="review"), follow=True)
        self.assertEqual(r.redirect_chain[-1][0], reverse("enrollment:review", args=[self.s1a.school_id]))
        self.assertContains(r, "ELEMENTARY TOTAL")
        self.assertEqual(r.context["totals"]["es"], (53, 48, 101))

    def test_editing_replaces_values_and_clearing_works(self):
        self.client.post(self.url, form_data(ATOK_ES))
        self.client.post(self.url, form_data({"kinder": (1, 2)}))
        t = compute_totals({e.level: (e.male, e.female) for e in EnrollmentEntry.objects.filter(record__school=self.s1a)})
        self.assertEqual(t["k12"], (1, 2, 3))

    def test_audit_entry_written_without_enrollment_values_or_secrets(self):
        self.client.post(self.url, form_data(ATOK_ES))
        entry = AuditLog.objects.get(action="record_saved")
        self.assertEqual(entry.user, self.u1)
        self.assertEqual(entry.district, self.d1)
        self.assertNotIn("13", str(entry.details.values()))
        self.assertNotIn(PW, str(entry.details))


class SubmissionTests(BaseTest):
    def setUp(self):
        super().setUp()
        self.login(self.u1)
        self.entry = reverse("enrollment:entry", args=[self.s1a.school_id])
        self.submit = reverse("enrollment:submit", args=[self.s1a.school_id])
        self.review = reverse("enrollment:review", args=[self.s1a.school_id])

    def _fill(self):
        self.client.post(self.entry, form_data(ATOK_ES))

    def test_cannot_submit_empty_or_incomplete(self):
        r = self.client.post(self.submit, {"confirm": "yes"})
        self.assertEqual(EnrollmentRecord.objects.get(school=self.s1a).status, RecordStatus.DRAFT)
        self.client.post(self.entry, form_data({"g1": (5, None)}))
        self.client.post(self.submit, {"confirm": "yes"})
        self.assertEqual(EnrollmentRecord.objects.get(school=self.s1a).status, RecordStatus.DRAFT)

    def test_confirmation_checkbox_required(self):
        self._fill()
        self.client.post(self.submit, {})
        self.assertEqual(EnrollmentRecord.objects.get(school=self.s1a).status, RecordStatus.DRAFT)

    def test_submit_locks_record(self):
        self._fill()
        self.client.post(self.submit, {"confirm": "yes"})
        rec = EnrollmentRecord.objects.get(school=self.s1a)
        self.assertEqual(rec.status, RecordStatus.SUBMITTED)
        self.assertEqual(rec.submitted_by, self.u1)
        self.assertIsNotNone(rec.submitted_at)
        r = self.client.post(self.entry, form_data({"g1": (1, 1)}))          # edit attempt
        self.assertRedirects(r, self.review)
        self.assertEqual(EnrollmentEntry.objects.get(record=rec, level="g1").male, 10)
        self.assertTrue(AuditLog.objects.filter(action="record_submitted").exists())

    def test_regular_user_cannot_reopen(self):
        self._fill()
        self.client.post(self.submit, {"confirm": "yes"})
        r = self.client.post(reverse("enrollment:reopen", args=[self.s1a.school_id]))
        self.assertEqual(r.status_code, 403)
        self.assertEqual(EnrollmentRecord.objects.get(school=self.s1a).status, RecordStatus.SUBMITTED)

    def test_district_admin_can_reopen_own_but_not_other_district(self):
        self._fill()
        self.client.post(self.submit, {"confirm": "yes"})
        self.login(self.admin1)
        self.client.post(reverse("enrollment:reopen", args=[self.s1a.school_id]))
        self.assertEqual(EnrollmentRecord.objects.get(school=self.s1a).status, RecordStatus.DRAFT)
        self.assertEqual(self.client.post(reverse("enrollment:reopen", args=[self.s2a.school_id])).status_code, 404)

    def test_reopen_requires_post(self):
        self.login(self.admin1)
        self.assertEqual(self.client.get(reverse("enrollment:reopen", args=[self.s1a.school_id])).status_code, 405)


class ManagementPermissionTests(BaseTest):
    def setUp(self):
        super().setUp()
        self.pending1 = make_user("pend1@deped.gov.ph", self.d1, approved=False)
        self.pending2 = make_user("pend2@deped.gov.ph", self.d2, approved=False)

    def act(self, user, action):
        return self.client.post(reverse("manage:user_action", args=[user.pk, action]))

    def test_regular_user_forbidden_from_manage_pages(self):
        self.login(self.u1)
        for name in ("manage:users", "manage:progress", "manage:audit"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 403, name)
        self.assertEqual(self.act(self.pending1, "approve").status_code, 403)
        self.pending1.refresh_from_db()
        self.assertEqual(self.pending1.approval_status, Approval.PENDING)

    def test_district_admin_approves_own_district_user(self):
        self.login(self.admin1)
        self.act(self.pending1, "approve")
        self.pending1.refresh_from_db()
        self.assertEqual(self.pending1.approval_status, Approval.APPROVED)
        self.assertEqual(self.pending1.approved_by, self.admin1)
        self.assertTrue(self.client.login(username=self.pending1.email, password=PW))

    def test_district_admin_cannot_touch_other_district_users(self):
        self.login(self.admin1)
        for action in ("approve", "reject", "suspend", "reactivate"):
            self.assertEqual(self.act(self.pending2, action).status_code, 404, action)
        self.pending2.refresh_from_db()
        self.assertEqual(self.pending2.approval_status, Approval.PENDING)
        self.assertTrue(self.pending2.is_active)

    def test_district_admin_cannot_modify_admins(self):
        self.login(self.admin1)
        other_admin = make_user("other.admin@deped.gov.ph", self.d1, "district_admin")
        self.assertEqual(self.act(other_admin, "suspend").status_code, 404)
        self.assertEqual(self.act(self.sysadmin, "suspend").status_code, 404)

    def test_cannot_act_on_self(self):
        self.login(self.sysadmin)
        self.assertEqual(self.act(self.sysadmin, "suspend").status_code, 403)

    def test_cannot_approve_unverified_email(self):
        u = make_user("unv@deped.gov.ph", self.d1, verified=False, approved=False)
        self.login(self.admin1)
        self.act(u, "approve")
        u.refresh_from_db()
        self.assertEqual(u.approval_status, Approval.PENDING)

    def test_suspend_blocks_login_and_reactivate_restores(self):
        self.login(self.admin1)
        self.act(self.u1, "suspend")
        self.assertFalse(self.client.login(username=self.u1.email, password=PW))
        self.login(self.admin1)
        self.act(self.u1, "reactivate")
        self.assertTrue(self.client.login(username=self.u1.email, password=PW))

    def test_actions_require_post_and_unknown_action_rejected(self):
        self.login(self.admin1)
        self.assertEqual(self.client.get(reverse("manage:user_action", args=[self.pending1.pk, "approve"])).status_code, 405)
        self.assertEqual(self.act(self.pending1, "delete").status_code, 403)

    def test_user_list_scoped_to_district(self):
        self.login(self.admin1)
        r = self.client.get(reverse("manage:users"), {"status": "pending"})
        self.assertContains(r, "pend1@deped.gov.ph")
        self.assertNotContains(r, "pend2@deped.gov.ph")
        self.login(self.sysadmin)
        r = self.client.get(reverse("manage:users"), {"status": "pending"})
        self.assertContains(r, "pend2@deped.gov.ph")

    def test_audit_log_scoped_and_actions_recorded(self):
        self.login(self.admin1)
        self.act(self.pending1, "approve")
        self.act_log = AuditLog.objects.get(action="user_approve")
        self.assertEqual(self.act_log.district, self.d1)
        AuditLog.objects.create(action="x_other", district=self.d2, actor_email="zz@deped.gov.ph")
        r = self.client.get(reverse("manage:audit"))
        self.assertContains(r, "user_approve")
        self.assertNotContains(r, "x_other")

    def test_progress_scoped(self):
        get_or_create_record(self.s1a, self.u1)
        self.login(self.admin1)
        r = self.client.get(reverse("manage:progress"))
        self.assertContains(r, "Atok")
        self.assertNotContains(r, "Bakun")



class ExportTests(BaseTest):
    def test_export_totals_and_layout(self):
        rec = get_or_create_record(self.s1a, self.u1)
        save_values(rec, ATOK_ES, self.u1)

        self.login(self.u1)
        response = self.client.get(reverse("enrollment:export"))

        self.assertIn("spreadsheetml", response["Content-Type"])

        ws = load_workbook(io.BytesIO(response.content)).active
        rows = list(ws.iter_rows(min_row=4, values_only=True))

        adoy = next(row for row in rows if row[2] == "Adoyunan ES")
        self.assertEqual(adoy[1], "100001")
        self.assertEqual(adoy[6:9], (6, 4, 10))  # Kinder M, F, T

        total_row = next(row for row in rows if row[2] == "TOTAL")

        # Locate the K-12 aggregate using the export schema rather
        # than assuming it occupies the final three worksheet columns.
        from .schema import EXPORT_ORDER

        enrollment_column = 6  # First six columns contain school details.
        k12_column = None

        for kind, key in EXPORT_ORDER:
            if kind == "agg" and key == "k12":
                k12_column = enrollment_column
                break
            enrollment_column += 3

        self.assertIsNotNone(k12_column, "K-12 aggregate is missing from EXPORT_ORDER")
        self.assertEqual(
            total_row[k12_column:k12_column + 3],
            (53, 48, 101),
        )

        self.assertEqual(ws.cell(2, 7).value, "Kinder")
